"""Local HTML Playwright simulation: save draft without final submit."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import pytest

from app.adapters.generic_portal import GenericPortalAdapter
from app.browser_session import browser_session
from app.models import AuthorConfig, AuthorEntry, PackageManifest
from app.package_builder import build_package
from app.settings import get_settings

PORTAL_HTML = """<!DOCTYPE html>
<html lang="en">
<head><meta charset="utf-8"><title>Example Journal Submit</title></head>
<body>
  <h1>Example Open Access Journal — Submission</h1>
  <a id="new-submission" href="#form">New Submission</a>
  <a id="my-drafts" href="#drafts">My Drafts</a>
  <div id="form">
    <label>Title <input id="manuscript-title" type="text"></label>
    <label>Abstract <textarea id="abstract"></textarea></label>
    <label>Keywords <input id="keywords" type="text"></label>
    <label>Corresponding email <input id="corresponding-email" type="email"></label>
    <label>Corresponding name <input id="corresponding-name" type="text"></label>
    <label>Country
      <select id="country">
        <option value="Brasil">Brasil</option>
        <option value="Other">Other</option>
      </select>
    </label>
    <label>Cover letter <textarea id="cover-letter"></textarea></label>
    <label>Blinded <input id="upload-blinded" type="file"></label>
    <label>Identified <input id="upload-identified" type="file"></label>
    <label>Title page <input id="upload-title-page" type="file"></label>
    <label>Cover file <input id="upload-cover" type="file"></label>
    <label>Checklist <input id="upload-checklist" type="file"></label>
    <button id="save-draft" type="button">Save Draft</button>
    <button id="final-submit" type="button">Final Submit</button>
    <button id="confirm-final-submit" type="button" style="display:none">Confirm Final</button>
    <div class="draft-saved-banner" id="draft-banner" style="display:none">Draft saved successfully</div>
    <div class="submission-success" id="success" style="display:none">Submitted</div>
    <div id="status"></div>
  </div>
  <script>
    let finalClicked = false;
    document.getElementById('save-draft').addEventListener('click', function() {
      document.getElementById('draft-banner').style.display = 'block';
      document.getElementById('status').textContent = 'DRAFT_SAVED';
    });
    document.getElementById('final-submit').addEventListener('click', function() {
      finalClicked = true;
      document.getElementById('confirm-final-submit').style.display = 'inline';
      document.getElementById('status').textContent = 'FINAL_CLICKED';
    });
    document.getElementById('confirm-final-submit').addEventListener('click', function() {
      document.getElementById('success').style.display = 'block';
      document.getElementById('status').textContent = 'SUBMITTED';
    });
    window.wasFinalClicked = function() { return finalClicked; };
  </script>
</body>
</html>
"""


@pytest.fixture
def portal_page(tmp_path: Path) -> Path:
    html = tmp_path / "portal.html"
    html.write_text(PORTAL_HTML, encoding="utf-8")
    return html


@pytest.fixture
def sample_package(tmp_path: Path) -> Path:
    # Minimal docx via python-docx
    from docx import Document

    src = tmp_path / "manuscript.docx"
    doc = Document()
    doc.add_heading("A Study on Safe Submission Bots", level=1)
    doc.add_paragraph("João Pedro Pereira Passos")
    doc.add_paragraph("Abstract")
    doc.add_paragraph("This is a sample abstract for testing package builds.")
    doc.add_paragraph("Keywords: safety, submission, dry-run")
    doc.add_paragraph("Introduction")
    doc.add_paragraph("Body text for the manuscript.")
    doc.save(str(src))

    out = tmp_path / "package"
    manifest = build_package(
        "MS-SIM",
        "example_journal",
        src,
        output_dir=out,
    )
    assert manifest.package_hash
    assert (out / "manifest.json").is_file()
    return out


@pytest.mark.asyncio
async def test_generic_adapter_saves_draft_without_final_submit(portal_page: Path, sample_package: Path):
    """Playwright local HTML: fill + save draft; final submit button never clicked."""
    get_settings.cache_clear()
    adapter = GenericPortalAdapter(selectors_key="example_journal")
    author = AuthorConfig(
        full_name="João Pedro Pereira Passos",
        publication_name="João Pedro Pereira Passos",
        corresponding_email="joaopedro.passos@mail.uft.edu.br",
        country="Brasil",
        authors=[
            AuthorEntry(
                full_name="João Pedro Pereira Passos",
                email="joaopedro.passos@mail.uft.edu.br",
                corresponding=True,
                order=1,
            )
        ],
    )
    manifest = PackageManifest.model_validate(
        json.loads((sample_package / "manifest.json").read_text(encoding="utf-8"))
    )
    url = portal_page.resolve().as_uri()

    async with browser_session(headless=True) as page:
        r_open = await adapter.open_submission(page, url)
        assert r_open.success

        r_fill = await adapter.fill_metadata(
            page,
            title="A Study on Safe Submission Bots",
            abstract="This is a sample abstract for testing package builds.",
            keywords="safety, submission, dry-run",
            author=author,
        )
        assert r_fill.success

        r_up = await adapter.upload_files(page, sample_package, manifest)
        assert r_up.success

        # Ensure final submit not clicked before draft
        final_before = await page.evaluate("window.wasFinalClicked()")
        assert final_before is False

        r_draft = await adapter.save_draft(page)
        assert r_draft.success
        assert r_draft.draft_saved is True
        assert r_draft.final_submitted is False

        status = await page.inner_text("#status")
        assert status == "DRAFT_SAVED"
        banner = await page.is_visible(".draft-saved-banner")
        assert banner is True

        # Critical: final submit must not have been clicked
        final_after = await page.evaluate("window.wasFinalClicked()")
        assert final_after is False
        success_visible = await page.is_visible("#success")
        assert success_visible is False

        # Adapter final_submit still refuses
        r_final = await adapter.final_submit(page)
        assert r_final.final_submitted is False
        # And page still not submitted via adapter path
        final_still = await page.evaluate("window.wasFinalClicked()")
        assert final_still is False


@pytest.mark.asyncio
async def test_fee_notice_blocks_draft(tmp_path: Path):
    html = tmp_path / "fee_portal.html"
    html.write_text(
        PORTAL_HTML.replace(
            "<h1>Example Open Access Journal — Submission</h1>",
            "<h1>Example</h1><div class='fee-notice'>Article Processing Charge APC payment required</div>",
        ),
        encoding="utf-8",
    )
    adapter = GenericPortalAdapter(
        selectors_key="example_journal",
        fee_keywords=["article processing charge", "apc", "payment required"],
    )
    async with browser_session(headless=True) as page:
        await page.goto(html.resolve().as_uri())
        result = await adapter.detect_blockers(page)
        assert result.blockers
        # save_draft should also stop
        r = await adapter.save_draft(page)
        assert r.success is False
        assert "FEE_OR_CONTRACT" in r.blockers


def test_package_builder_creates_expected_files(sample_package: Path):
    names = {p.name for p in sample_package.iterdir()}
    assert "manifest.json" in names
    assert "requirements_snapshot.json" in names
    assert any("blinded" in n for n in names)
    assert any("identified" in n for n in names)
    assert any("title_page" in n for n in names)
    assert any("cover_letter" in n for n in names)
    assert any("checklist" in n for n in names)
    data = json.loads((sample_package / "manifest.json").read_text(encoding="utf-8"))
    assert data["corresponding_email"] == "joaopedro.passos@mail.uft.edu.br"
    assert data["author_order"] == ["João Pedro Pereira Passos"]
